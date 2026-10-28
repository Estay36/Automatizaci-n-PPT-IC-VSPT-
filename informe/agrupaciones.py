"""Lectura de la hoja "Agrupaciones" de la planilla *Categorías Vinos*.

La hoja tiene tres bloques, cada uno encabezado por un rótulo en la columna A:

* ``COMPAÑÍA``: una columna por compañía (VSPT, CONCHA Y TORO, ...) con sus
  marcas; la columna "OTRAS VIÑAS" lista *empresas* (no marcas).
* ``CATEGORÍA``: una columna por categoría (MASIVOS, BOTELLAS FINAS, ...) con sus
  marcas. Una celda como "OTRAS BOTELLAS FINAS*" dentro de una lista remite a la
  columna homónima, cuyas marcas se incorporan a esa categoría.
* ``SUB CATEGORÍA``: listas BLYS VINOS y BLYS TOTAL. Lo que aparece bajo
  "Tentativo incluir a futuro" queda fuera hasta que se mueva a la lista.

El parser localiza los bloques por su rótulo, por lo que tolera filas o
columnas agregadas mientras se mantenga esa estructura.
"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import BinaryIO

import openpyxl

from . import config

HOJA = "Agrupaciones"


def clave(texto: object) -> str:
    """Normaliza un nombre para compararlo entre la hoja y la base.

    Mayúsculas, sin tildes, sin paréntesis ni asteriscos y sin signos:
    ``"Royal Guard (Sabores)*"`` -> ``"ROYALGUARD"``, ``"B-MIX"`` -> ``"BMIX"``.
    """
    if texto is None:
        return ""
    s = str(texto).upper()
    s = re.sub(r"\(.*?\)", "", s)
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^A-Z0-9Ñ]", "", s)


def _categoria_canonica(encabezado: str) -> str | None:
    k = clave(encabezado)
    if k.startswith(("OTRA", "OTRO")):
        return None
    if "MASIV" in k:
        return "MASIVOS"
    if "BOTELLA" in k:
        return "BOTELLAS FINAS"
    if "ESPUM" in k or "ESUM" in k:  # la hoja trae "ESUMANTES"
        return "ESPUMANTES"
    if "COOLER" in k or "SANGR" in k:
        return "COOLERS Y SANGRÍA"
    return None


@dataclass
class Agrupaciones:
    """Listas de la hoja Agrupaciones, ya normalizadas con :func:`clave`."""

    # compañía -> claves de marca
    marcas_compania: dict[str, set[str]] = field(default_factory=dict)
    # claves de empresa que caen en "OTRAS VIÑAS"
    empresas_otras: set[str] = field(default_factory=set)
    # categoría -> claves de marca
    marcas_categoria: dict[str, set[str]] = field(default_factory=dict)
    blys_vinos: set[str] = field(default_factory=set)
    blys_total: set[str] = field(default_factory=set)  # incluye blys_vinos
    blys_futuro: set[str] = field(default_factory=set)  # excluidas por ahora
    comentarios: list[str] = field(default_factory=list)
    companias: list[str] = field(default_factory=list)  # en orden de la hoja

    def categorias_de(self, clave_marca: str) -> list[str]:
        return [c for c, s in self.marcas_categoria.items() if clave_marca in s]

    def compania_de_marca(self, clave_marca: str) -> str | None:
        for comp, marcas in self.marcas_compania.items():
            if clave_marca in marcas:
                return comp
        return None

    def resumen(self) -> dict[str, int]:
        r = {f"Compañía · {c}": len(m) for c, m in self.marcas_compania.items()}
        r["Compañía · OTRAS VIÑAS (empresas)"] = len(self.empresas_otras)
        r.update({f"Categoría · {c}": len(m) for c, m in self.marcas_categoria.items()})
        r["BLYS VINOS"] = len(self.blys_vinos)
        r["BLYS TOTAL"] = len(self.blys_total)
        return r


def _valores(ws, fila: int, col: int, hasta: int) -> list[str]:
    out = []
    for r in range(fila, hasta + 1):
        v = ws.cell(r, col).value
        if v is not None and str(v).strip():
            out.append(str(v).strip())
    return out


def _filas_bloque(ws) -> dict[str, int]:
    filas = {}
    for r in range(1, ws.max_row + 1):
        k = clave(ws.cell(r, 1).value)
        if k in ("COMPANIA", "CATEGORIA", "SUBCATEGORIA") and k not in filas:
            filas[k] = r
    faltan = {"COMPANIA", "CATEGORIA", "SUBCATEGORIA"} - set(filas)
    if faltan:
        raise ValueError(f"La hoja '{HOJA}' no tiene los bloques: {sorted(faltan)}")
    return filas


def _fila_encabezado(ws, fila_rotulo: int) -> int:
    """Primera fila no vacía bajo el rótulo del bloque (puede haber filas en blanco)."""
    for r in range(fila_rotulo + 1, ws.max_row + 1):
        if any(ws.cell(r, c).value not in (None, "") for c in range(1, ws.max_column + 1)):
            return r
    raise ValueError(f"Bloque sin encabezados en la fila {fila_rotulo} de '{HOJA}'.")


def leer_agrupaciones(fuente: str | Path | BinaryIO) -> Agrupaciones:
    wb = openpyxl.load_workbook(fuente, data_only=True)
    if HOJA not in wb.sheetnames:
        raise ValueError(f"La planilla no tiene la hoja '{HOJA}'.")
    ws = wb[HOJA]
    filas = _filas_bloque(ws)
    ag = Agrupaciones()
    ultima = ws.max_row
    alias = {clave(k): clave(v) for k, v in config.ALIAS_MARCAS.items()}

    def norm(nombres: list[str]) -> set[str]:
        return {alias.get(clave(n), clave(n)) for n in nombres}

    # --- Compañías --------------------------------------------------------- #
    enc, fin = _fila_encabezado(ws, filas["COMPANIA"]), filas["CATEGORIA"] - 1
    for col in range(1, ws.max_column + 1):
        titulo = ws.cell(enc, col).value
        if not titulo:
            continue
        titulo = str(titulo).strip().upper()
        items = _valores(ws, enc + 1, col, fin)
        if clave(titulo).startswith("OTRA"):
            ag.empresas_otras = norm(items)
        else:
            ag.marcas_compania[titulo] = norm(items) | {clave(titulo)}
            ag.companias.append(titulo)

    # --- Categorías -------------------------------------------------------- #
    enc, fin = _fila_encabezado(ws, filas["CATEGORIA"]), filas["SUBCATEGORIA"] - 1
    listas: dict[str, list[str]] = {}
    for col in range(1, ws.max_column + 1):
        titulo = ws.cell(enc, col).value
        if not titulo:
            continue
        titulo = str(titulo).strip()
        items = _valores(ws, enc + 1, col, fin)
        if clave(titulo).startswith("COMENTARIO"):
            ag.comentarios += items
        else:
            listas[clave(titulo)] = items
    for k_titulo, items in listas.items():
        cat = _categoria_canonica(k_titulo)
        if cat is None:
            continue
        marcas: list[str] = []
        for item in items:
            k = clave(item)
            if k in listas and _categoria_canonica(k) is None:  # remite a otra lista
                marcas += listas[k]
            else:
                marcas.append(item)
        ag.marcas_categoria[cat] = norm(marcas)
    ag.marcas_categoria = {c: ag.marcas_categoria.get(c, set())
                           for c in config.ORDEN_CATEGORIAS}

    # --- BLYS ---------------------------------------------------------------- #
    enc = _fila_encabezado(ws, filas["SUBCATEGORIA"])
    for col in range(1, ws.max_column + 1):
        titulo = clave(ws.cell(enc, col).value)
        if titulo.startswith("COMENTARIO"):
            ag.comentarios += _valores(ws, enc + 1, col, ultima)
            continue
        if titulo not in ("BLYSVINOS", "BLYSTOTAL"):
            continue
        actuales, futuro, en_futuro = [], [], False
        for r in range(enc + 1, ultima + 1):
            v = ws.cell(r, col).value
            if v is None or not str(v).strip():
                continue
            if "FUTURO" in clave(v):
                en_futuro = True
                continue
            (futuro if en_futuro else actuales).append(str(v).strip())
        if titulo == "BLYSVINOS":
            ag.blys_vinos = norm(actuales)
            ag.blys_futuro |= norm(futuro)
        else:
            ag.blys_total = norm(actuales)
            ag.blys_futuro |= norm(futuro)
    ag.blys_total |= ag.blys_vinos
    return ag
