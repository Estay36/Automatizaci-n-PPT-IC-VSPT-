"""Primitivas de PowerPoint: diapositivas desde la plantilla, gráficos nativos y tablas.

Los gráficos son nativos (editables en PowerPoint, con su planilla de datos) y
replican el estilo del informe de referencia: columnas apiladas con gap 45 %,
etiquetas internas, totales sobre cada columna y paleta fija por entidad.
"""
from __future__ import annotations

import math
from dataclasses import dataclass
from typing import Sequence

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from . import config

# Geometría (EMU) de la zona de contenido, alineada con el informe de referencia.
ANCHO_SLIDE, ALTO_SLIDE = 12192000, 6858000
X0, Y0 = 442914, 1640000
ANCHO, ALTO = 11306175, 4780000
Y_FUENTE = 6581001  # parte superior de la línea de fuente en la referencia

GRIS_TEXTO = "595959"
GRIS_GRILLA = "D9D9D9"
GRIS_ENCABEZADO = "D9D9D9"
GRIS_TOTAL = "F2F2F2"
ROJO_NEGATIVO = "C00000"
TABLA_SIN_ESTILO = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # "No Style, Table Grid"


# --------------------------------------------------------------------------- #
# Colores
# --------------------------------------------------------------------------- #
def texto_sobre(hex_fondo: str) -> str:
    """Blanco sobre colores oscuros, gris oscuro sobre claros (luminancia WCAG)."""
    r, g, b = (int(hex_fondo[i:i + 2], 16) / 255 for i in (0, 2, 4))
    lin = [c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4 for c in (r, g, b)]
    lum = 0.2126 * lin[0] + 0.7152 * lin[1] + 0.0722 * lin[2]
    return "FFFFFF" if lum < 0.35 else "262626"


class Colores:
    """Asigna colores estables por entidad y sin repetir dentro de un gráfico."""

    def __init__(self, fijos: dict[str, str], respaldo: list[str]):
        self.fijos = fijos
        self.respaldo = respaldo
        self._asignados: dict[str, str] = {}

    def para(self, entidades: Sequence[str]) -> dict[str, str]:
        usados: set[str] = set()
        reservados = {self.fijos[e.upper()] for e in entidades if e.upper() in self.fijos}
        out = {}
        for e in entidades:
            k = e.upper()
            c = self.fijos.get(k) or self._asignados.get(k)
            if c is None or c in usados:
                libres = [p for p in self.respaldo if p not in usados and p not in reservados]
                libres = libres or [p for p in self.respaldo if p not in usados] or self.respaldo
                c = libres[0]
                self._asignados.setdefault(k, c)
            usados.add(c)
            out[e] = c
        return out


# --------------------------------------------------------------------------- #
# Formato de números (es-CL)
# --------------------------------------------------------------------------- #
def num(x: float | None, dec: int = 0) -> str:
    if x is None or (isinstance(x, float) and math.isnan(x)):
        return ""
    s = f"{x:,.{dec}f}"
    return s.replace(",", "§").replace(".", ",").replace("§", ".")


def pct(x: float | None, dec: int = 1, tope: float | None = None) -> str:
    """Porcentaje es-CL. Con ``tope`` (p. ej. 9.99), valores mayores se muestran
    como ">999%" para no inflar la tabla con variaciones sobre bases ínfimas."""
    if x is None or (isinstance(x, float) and math.isnan(x)):
        return ""
    if tope is not None and x > tope:
        return ">" + num(tope * 100, 0) + "%"
    return num(x * 100, dec) + "%"


def _limpio(v) -> float | None:
    if v is None:
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return None if math.isnan(f) else f


# --------------------------------------------------------------------------- #
# Presentación y diapositivas
# --------------------------------------------------------------------------- #
class Deck:
    def __init__(self, plantilla=config.PLANTILLA):
        self.prs = Presentation(str(plantilla))
        self._disenos = {l.name: l for m in self.prs.slide_masters for l in m.slide_layouts}

    def nueva(self, diseno: str):
        if diseno not in self._disenos:
            raise KeyError(f"La plantilla no tiene el diseño '{diseno}'.")
        return self.prs.slides.add_slide(self._disenos[diseno])

    def guardar(self, destino) -> None:
        self.prs.save(destino)


def _run(parrafo, texto: str, *, tam: float | None = None, negrita: bool | None = None,
         color: str | None = None, fuente: str | None = None):
    r = parrafo.add_run()
    r.text = texto
    f = r.font
    if tam:
        f.size = Pt(tam)
    if negrita is not None:
        f.bold = negrita
    if color:
        f.color.rgb = RGBColor.from_string(color)
    if fuente:
        f.name = fuente
    return r


def poner_titulo(slide, titulo: str, subtitulo: str | None = None) -> None:
    """Título en el marcador del diseño + subtítulo (periodo) en la misma caja."""
    tf = slide.shapes.title.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    _run(p, titulo, tam=34)
    if subtitulo:
        p2 = tf.add_paragraph()
        _run(p2, subtitulo, tam=16)


def caja_texto(slide, x, y, w, h, texto: str, *, tam=12, negrita=False, color=None,
               alinear=PP_ALIGN.LEFT, fuente=None, anclaje=MSO_ANCHOR.TOP, nombre=None):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    if nombre:
        tb.name = nombre
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anclaje
    for i, linea in enumerate(texto.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinear
        _run(p, linea, tam=tam, negrita=negrita, color=color, fuente=fuente)
    return tb


def poner_fuente(slide, texto: str) -> None:
    """Nota de fuente abajo a la derecha, como en el informe de referencia."""
    # Empieza a la derecha del pie "Confidential..." del patrón para no solaparlo.
    x = 7300000
    tb = caja_texto(slide, x, Y_FUENTE - 130000, ANCHO_SLIDE - x - 80000, 400000,
                    texto, tam=9, alinear=PP_ALIGN.RIGHT, anclaje=MSO_ANCHOR.BOTTOM,
                    nombre="Fuente")
    tb.text_frame.margin_top = tb.text_frame.margin_bottom = 0


# --------------------------------------------------------------------------- #
# Helpers XML de gráficos
# --------------------------------------------------------------------------- #
def _c(tag: str, **attrs):
    el = etree.Element(qn(f"c:{tag}"))
    for k, v in attrs.items():
        el.set(k, str(v))
    return el


def layout_manual(chart, x: float, y: float, w: float, h: float) -> None:
    """Fija el área de trazado (fracciones del marco) para alinear gráficos."""
    plot_area = chart._chartSpace.chart.plotArea  # noqa: SLF001
    viejo = plot_area.find(qn("c:layout"))
    if viejo is not None:
        plot_area.remove(viejo)
    layout = _c("layout")
    ml = etree.SubElement(layout, qn("c:manualLayout"))
    for tag, val in (("layoutTarget", "inner"), ("xMode", "edge"), ("yMode", "edge"),
                     ("x", x), ("y", y), ("w", w), ("h", h)):
        etree.SubElement(ml, qn(f"c:{tag}")).set("val", str(val))
    plot_area.insert(0, layout)


def ocultar_en_leyenda(chart, idx_series: Sequence[int]) -> None:
    legend = chart._chartSpace.chart.find(qn("c:legend"))  # noqa: SLF001
    if legend is None:
        return
    pos = legend.find(qn("c:legendPos"))
    for i in idx_series:
        le = _c("legendEntry")
        etree.SubElement(le, qn("c:idx")).set("val", str(i))
        etree.SubElement(le, qn("c:delete")).set("val", "1")
        pos.addnext(le)


def ocultar_etiqueta(serie, idx: int) -> None:
    dlbls = serie._element.get_or_add_dLbls()  # noqa: SLF001
    dlbl = dlbls.get_or_add_dLbl_for_point(idx)
    for hijo in list(dlbl):
        if hijo.tag != qn("c:idx"):
            dlbl.remove(hijo)
    etree.SubElement(dlbl, qn("c:delete")).set("val", "1")


def _rotar_texto_eje(eje, grados: int) -> None:
    txpr = eje._element.get_or_add_txPr()  # noqa: SLF001
    txpr.find(qn("a:bodyPr")).set("rot", str(int(grados * 60000)))


def _estilo_ejes(chart, *, formato_valor: str, max_valor: float | None,
                 rot_categorias: int, titulo_valor: str | None, tam: float) -> None:
    va, ca = chart.value_axis, chart.category_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor.from_string(GRIS_GRILLA)
    va.major_gridlines.format.line.width = Pt(0.5)
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(tam)
    va.tick_labels.font.color.rgb = RGBColor.from_string(GRIS_TEXTO)
    va.tick_labels.number_format = formato_valor
    va.tick_labels.number_format_is_linked = False
    va.minimum_scale = 0
    if max_valor is not None:
        va.maximum_scale = max_valor
    if titulo_valor:
        va.has_title = True
        va.axis_title.text_frame.text = titulo_valor
        f = va.axis_title.text_frame.paragraphs[0].runs[0].font
        f.size, f.bold = Pt(tam), False
        f.color.rgb = RGBColor.from_string(GRIS_TEXTO)
    ca.format.line.color.rgb = RGBColor.from_string("A6A6A6")
    ca.tick_labels.font.size = Pt(tam)
    ca.tick_labels.font.color.rgb = RGBColor.from_string(GRIS_TEXTO)
    ca.has_major_gridlines = False
    if rot_categorias:
        _rotar_texto_eje(ca, rot_categorias)


# --------------------------------------------------------------------------- #
# Gráfico de columnas apiladas (absoluto o 100 %)
# --------------------------------------------------------------------------- #
def columnas_apiladas(
    slide, x, y, w, h,
    categorias: Sequence[str],
    series: Sequence[tuple[str, Sequence[float | None]]],
    colores: dict[str, str],
    *,
    porcentaje: bool = False,
    totales: Sequence[str] | None = None,
    formato_etiqueta: str = "#,##0",
    formato_valor: str = "#,##0",
    umbral_etiqueta: float = 0.04,
    leyenda: bool = True,
    trazado: tuple[float, float, float, float] | None = None,
    titulo_valor: str | None = None,
    rot_categorias: int = 0,
    rot_totales: int = 0,
    tam: float = 9,
    tam_etiquetas: float | None = None,
    gap: int = 45,
):
    """Columnas apiladas con etiquetas internas y totales opcionales arriba.

    ``series`` son valores ya escalados. Si ``porcentaje`` es True, deben ser
    participaciones (0-1) y el gráfico es 100 % apilado. Las etiquetas de segmentos
    que pesan menos que ``umbral_etiqueta`` dentro de su columna se ocultan.
    ``totales`` es el texto a mostrar sobre cada columna ("" para ninguno).
    """
    tam_etiquetas = tam_etiquetas or tam
    cd = CategoryChartData(number_format=formato_etiqueta)
    cd.categories = list(categorias)
    limpias = [(n, [_limpio(v) for v in vals]) for n, vals in series]
    for nombre, vals in limpias:
        cd.add_series(nombre, vals)
    con_total = totales is not None
    if con_total:
        cd.add_series("Total", [0.0] * len(categorias))
    tipo = XL_CHART_TYPE.COLUMN_STACKED_100 if porcentaje else XL_CHART_TYPE.COLUMN_STACKED
    gf = slide.shapes.add_chart(tipo, Emu(x), Emu(y), Emu(w), Emu(h), cd)
    chart = gf.chart
    chart.font.size = Pt(tam)
    plot = chart.plots[0]
    plot.gap_width = gap
    plot.overlap = 100

    # Participación de cada segmento dentro de su columna (para ocultar etiquetas).
    sumas = [sum(v[i] or 0 for _, v in limpias) for i in range(len(categorias))]
    maximo = max(sumas, default=0)
    for k, (nombre, vals) in enumerate(limpias):
        s = plot.series[k]
        color = colores.get(nombre, "A5A5A5")
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = RGBColor.from_string(color)
        s.format.line.fill.background()
        s.invert_if_negative = False
        dl = s.data_labels
        dl.show_value = True
        dl.number_format = formato_etiqueta
        dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.CENTER
        dl.font.size = Pt(tam_etiquetas)
        dl.font.color.rgb = RGBColor.from_string(texto_sobre(color))
        for i, v in enumerate(vals):
            # Se oculta si el segmento pesa poco en su columna o, en gráficos
            # absolutos, si es muy bajo respecto de la columna más alta (no cabe).
            chico = (not porcentaje and v is not None and maximo > 0
                     and v / maximo < umbral_etiqueta)
            if v is None or sumas[i] <= 0 or (v / sumas[i]) < umbral_etiqueta or chico:
                ocultar_etiqueta(s, i)

    if con_total:
        s = plot.series[len(limpias)]
        s.format.fill.background()
        s.format.line.fill.background()
        for i, texto in enumerate(totales):
            p_dl = s.points[i].data_label
            if not texto:
                ocultar_etiqueta(s, i)
                continue
            p_dl.position = XL_LABEL_POSITION.INSIDE_BASE
            tf = p_dl.text_frame
            tf.text = texto
            f = tf.paragraphs[0].runs[0].font
            f.size, f.bold = Pt(tam), True
            f.color.rgb = RGBColor.from_string("404040")
            if rot_totales:
                body = p_dl._dLbl.find(qn("c:tx")).find(qn("c:rich")).find(qn("a:bodyPr"))  # noqa: SLF001
                body.set("rot", str(int(rot_totales * 60000)))
                body.set("vert", "horz")

    _estilo_ejes(chart, formato_valor=formato_valor, max_valor=1 if porcentaje else None,
                 rot_categorias=rot_categorias, titulo_valor=titulo_valor, tam=tam)
    chart.has_legend = leyenda
    if leyenda:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(tam)
        if con_total:
            ocultar_en_leyenda(chart, [len(limpias)])
    if trazado:
        layout_manual(chart, *trazado)
    return chart


# --------------------------------------------------------------------------- #
# Dona doble (anillo interior SOS, exterior SOV)
# --------------------------------------------------------------------------- #
def dona_doble(slide, x, y, w, h, entidades: Sequence[str],
               interior: Sequence[float], exterior: Sequence[float],
               colores: dict[str, str], nombres=("SOS", "SOV"), umbral=0.03, tam=11):
    cd = CategoryChartData(number_format="0%")
    cd.categories = list(entidades)
    cd.add_series(nombres[0], [float(v) for v in interior])
    cd.add_series(nombres[1], [float(v) for v in exterior])
    gf = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Emu(x), Emu(y), Emu(w), Emu(h), cd)
    chart = gf.chart
    plot = chart.plots[0]
    plot.vary_by_categories = True
    dona = chart._chartSpace.chart.plotArea.find(qn("c:doughnutChart"))  # noqa: SLF001
    hole = dona.find(qn("c:holeSize"))
    if hole is None:
        hole = etree.SubElement(dona, qn("c:holeSize"))
    hole.set("val", "35")
    for s, vals in zip(plot.series, (interior, exterior)):
        s.data_labels.show_value = True
        s.data_labels.number_format = "0%"
        s.data_labels.number_format_is_linked = False
        s.data_labels.font.size = Pt(tam)
        for i, e in enumerate(entidades):
            color = colores[e]
            pt = s.points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = RGBColor.from_string(color)
            pt.format.line.color.rgb = RGBColor.from_string("FFFFFF")
            if vals[i] < umbral:
                ocultar_etiqueta(s, i)
            else:
                pt.data_label.font.size = Pt(tam)
                pt.data_label.font.color.rgb = RGBColor.from_string(texto_sobre(color))
                pt.data_label.number_format = "0%"
                pt.data_label.number_format_is_linked = False
                pt.data_label.show_value = True
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.RIGHT
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(tam)
    return chart


# --------------------------------------------------------------------------- #
# Tablas
# --------------------------------------------------------------------------- #
def _celda(celda, texto: str, *, tam=8, negrita=False, fondo=None, alinear=PP_ALIGN.RIGHT,
           color=None):
    celda.margin_left = celda.margin_right = Emu(36000)
    celda.margin_top = celda.margin_bottom = Emu(0)
    celda.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = celda.text_frame
    tf.clear()
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = alinear
    p.space_before = p.space_after = Pt(0)
    p.line_spacing = 1.0
    if texto:
        _run(p, texto, tam=tam, negrita=negrita, color=color or "000000")
    # La marca de fin de párrafo define el alto mínimo de celdas vacías: sin esto
    # PowerPoint usa 18 pt y la tabla crece por sobre lo previsto.
    p._p.get_or_add_endParaRPr().set("sz", str(int(tam * 100)))  # noqa: SLF001
    if fondo:
        celda.fill.solid()
        celda.fill.fore_color.rgb = RGBColor.from_string(fondo)
    else:
        celda.fill.background()


def nueva_tabla(slide, filas: int, anchos: Sequence[int], x: int, y: int, alto_fila: int):
    shape = slide.shapes.add_table(filas, len(anchos), Emu(x), Emu(y), Emu(sum(anchos)),
                                   Emu(alto_fila * filas))
    tabla = shape.table
    tbl_pr = shape._element.graphic.graphicData.tbl.tblPr  # noqa: SLF001
    for attr in ("firstRow", "bandRow"):
        tbl_pr.set(attr, "0")
    estilo = tbl_pr.find(qn("a:tableStyleId"))
    if estilo is None:
        estilo = etree.SubElement(tbl_pr, qn("a:tableStyleId"))
    estilo.text = TABLA_SIN_ESTILO
    for i, a in enumerate(anchos):
        tabla.columns[i].width = Emu(a)
    for r in tabla.rows:
        r.height = Emu(alto_fila)
    return tabla


def _sin_bordes(celda) -> None:
    """Quita los bordes propios de la celda (filas separadoras)."""
    tc_pr = celda._tc.get_or_add_tcPr()  # noqa: SLF001
    for lado in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        viejo = tc_pr.find(qn(lado))
        if viejo is not None:
            tc_pr.remove(viejo)
    for pos, lado in enumerate(("a:lnL", "a:lnR")):
        ln = etree.Element(qn(lado), w="0")
        etree.SubElement(ln, qn("a:noFill"))
        tc_pr.insert(pos, ln)


@dataclass
class BloqueTabla:
    """Un año de la tabla mensual: (nombre, 12 meses, total, share) por fila."""

    anio: int
    filas: Sequence[tuple[str, Sequence[str], str, str]]
    total: tuple[Sequence[str], str, str]


def tabla_mensual(slide, x: int, y: int, anchos: Sequence[int], alto_fila: int, *,
                  bloques: Sequence[BloqueTabla], etiqueta_share: str,
                  variacion: Sequence[float | None], variacion_total: float | None,
                  alto_separador: int = 70000, tam: float = 8) -> None:
    """Tabla "MARCA | Ene..Dic | TOTAL | SOS" por año + fila de variación.

    Todo va en **una sola tabla** (con filas separadoras sin bordes) para que el
    alto real de las filas, que decide PowerPoint según el texto, nunca provoque
    solapes entre bloques.
    """
    n = sum(3 + len(b.filas) for b in bloques) + len(bloques)  # + separadores
    n += 1  # variación
    t = nueva_tabla(slide, n, anchos, x, y, alto_fila)
    ultima = len(anchos) - 1
    enc = dict(tam=tam, negrita=True, fondo=GRIS_ENCABEZADO)
    r = 0
    for b in bloques:
        _celda(t.cell(r, 0), "MARCA", alinear=PP_ALIGN.LEFT, **enc)
        _celda(t.cell(r + 1, 0), "", **enc)
        t.cell(r, 0).merge(t.cell(r + 1, 0))
        _celda(t.cell(r, 1), str(b.anio), alinear=PP_ALIGN.CENTER, **enc)
        for m in range(2, 13):
            _celda(t.cell(r, m), "", **enc)
        t.cell(r, 1).merge(t.cell(r, 12))
        for lbl, col in (("TOTAL", ultima - 1), (etiqueta_share, ultima)):
            _celda(t.cell(r, col), lbl, alinear=PP_ALIGN.CENTER, **enc)
            _celda(t.cell(r + 1, col), "", **enc)
            t.cell(r, col).merge(t.cell(r + 1, col))
        for m, mes in enumerate(config.MESES_CORTOS, start=1):
            _celda(t.cell(r + 1, m), mes, alinear=PP_ALIGN.CENTER, **enc)
        r += 2
        for nombre, meses, tot, share in b.filas:
            _celda(t.cell(r, 0), nombre, tam=tam, alinear=PP_ALIGN.LEFT)
            for m, v in enumerate(meses, start=1):
                _celda(t.cell(r, m), v, tam=tam)
            _celda(t.cell(r, ultima - 1), tot, tam=tam, negrita=True, fondo=GRIS_TOTAL)
            _celda(t.cell(r, ultima), share, tam=tam, fondo=GRIS_TOTAL)
            r += 1
        _celda(t.cell(r, 0), "TOTAL", alinear=PP_ALIGN.LEFT, **enc)
        for m, v in enumerate(b.total[0], start=1):
            _celda(t.cell(r, m), v, **enc)
        _celda(t.cell(r, ultima - 1), b.total[1], **enc)
        _celda(t.cell(r, ultima), b.total[2], **enc)
        r += 1
        # separador
        t.rows[r].height = Emu(alto_separador)
        for c in range(len(anchos)):
            _celda(t.cell(r, c), "", tam=4)
            _sin_bordes(t.cell(r, c))
        r += 1
    _celda(t.cell(r, 0), "Var.% año anterior", alinear=PP_ALIGN.LEFT, **enc)
    for m, v in enumerate(list(variacion) + [variacion_total], start=1):
        color = ROJO_NEGATIVO if (v is not None and v < 0) else None
        _celda(t.cell(r, m), pct(v, 0, tope=9.99) if v is not None else "", tam=tam, negrita=True,
               color=color)
    _celda(t.cell(r, ultima), "", tam=tam)
