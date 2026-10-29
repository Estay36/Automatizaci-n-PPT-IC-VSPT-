"""Constructores de cada tipo de diapositiva del informe (hoja "Informe").

Cada función recibe el ``Deck``, el contexto común (datos, periodos, colores,
opciones) y la sección, y agrega una diapositiva. Si no hay datos suficientes,
no agrega nada y deja un aviso en ``ctx.avisos``.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date
from typing import Callable

import pandas as pd
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

from . import config, metricas
from .datos import Periodos
from .pptx_utils import (ALTO, ANCHO, X0, Y0, Colores, Deck, caja_texto,
                         columnas_apiladas, dona_doble, num, pct, poner_fuente,
                         poner_titulo)
from .pptx_utils import BloqueTabla
from .pptx_utils import tabla_mensual as dibujar_tabla_mensual
from .trazabilidad import Bloque, Entrada, Trazabilidad, Variacion


# --------------------------------------------------------------------------- #
# Estructuras
# --------------------------------------------------------------------------- #
@dataclass
class Seccion:
    clave: str                  # identificador interno
    nombre: str                 # texto del separador y prefijo de subtítulos
    entidad: str                # "COMPAÑÍA", "CATEGORÍA" o "MARCA" (para títulos)
    dim: str                    # columna de agrupación
    filtro: Callable[[pd.DataFrame], pd.Series]
    diapositivas: tuple[str, ...]
    target: str = config.TARGET_DEFECTO
    orden_fijo: list[str] | None = None
    separador: str | None = None  # texto del separador (por defecto ``nombre``)
    universo: str = ""            # filtro en palabras, para la trazabilidad


@dataclass
class Opciones:
    fecha_corte: date
    anio_inicio: int | None = None
    top_grafico: int = 10
    top_tabla: int = 10
    top_mix: int = 10
    top_sov: int = 8
    titulo: str = "INFORME DE COMPETENCIA – VSPT"


@dataclass
class Contexto:
    df: pd.DataFrame
    per: Periodos
    colores: Colores
    op: Opciones
    avisos: list[str] = field(default_factory=list)
    traza: Trazabilidad | None = None


# --------------------------------------------------------------------------- #
# Trazabilidad
# --------------------------------------------------------------------------- #
def _registrar(ctx: Contexto, deck: Deck, sec: Seccion, tipo: str, titulo: str,
               subtitulo: str, **kw) -> None:
    """Deja constancia de los datos de la diapositiva recién creada."""
    if ctx.traza is None:
        return
    ctx.traza.registrar(Entrada(n=len(deck.prs.slides), seccion=sec.nombre, tipo=tipo,
                                titulo=titulo, subtitulo=subtitulo, universo=sec.universo,
                                **kw))


def _asignacion(entidades, visibles) -> dict[str, str]:
    visibles = set(visibles)
    return {e: (e if e in visibles else config.OTRAS) for e in entidades}


def _regla_series(sec: Seccion, n: int, criterio: str) -> str:
    if sec.orden_fijo is not None:
        return "Series: todas las entidades, en orden fijo."
    return f"Series: top {n} {criterio}; el resto se suma en OTRAS."


def _texto_periodos(per: Periodos) -> str:
    return (f"TT AAAA: Año = AAAA. YTD {per.anio_anterior}: Año = {per.anio_anterior} y "
            f"ytd = VERDADERO. YTD {per.anio_actual}: Año = {per.anio_actual}.")


# --------------------------------------------------------------------------- #
# Textos de fechas
# --------------------------------------------------------------------------- #
def _fecha(d: date, con_anio: bool = True) -> str:
    s = f"{d.day:02d} DE {config.MESES[d.month - 1]}"
    return f"{s} {d.year}" if con_anio else s


def rango(per: Periodos, desde_anio: int | None = None) -> str:
    """"01 DE ENERO AL 02 DE OCTUBRE 2026" o "01 DE ENERO 2025 AL ..." ."""
    corte = _fecha(per.fecha_corte)
    if desde_anio is None or desde_anio == per.anio_actual:
        return f"01 DE ENERO AL {corte}"
    return f"01 DE ENERO {desde_anio} AL {corte}"


def _sub(sec: Seccion, texto: str) -> str:
    return f"{sec.nombre}  |  {texto}"


def _fuente_trp(sec: Seccion) -> str:
    return config.FUENTE_TRP.format(target=config.TARGETS_TRP.get(sec.target, sec.target))


def _limpiar_marcadores(slide, conservar_titulo: bool) -> None:
    for ph in list(slide.placeholders):
        if conservar_titulo and ph.placeholder_format.type == 1:  # TITLE
            continue
        ph._element.getparent().remove(ph._element)  # noqa: SLF001


# --------------------------------------------------------------------------- #
# Portada y separadores
# --------------------------------------------------------------------------- #
def portada(deck: Deck, ctx: Contexto) -> None:
    s = deck.nueva(config.DISENO_PORTADA)
    _limpiar_marcadores(s, conservar_titulo=False)
    tb = caja_texto(s, 556727, 1974255, 11078546, 1200329, ctx.op.titulo, tam=48,
                    color=config.COLOR_MARCA_TEXTO, alinear=PP_ALIGN.CENTER,
                    fuente="Futura Std Book", nombre="Título portada")
    p = tb.text_frame.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = rango(ctx.per)
    r.font.size = Pt(24)
    r.font.name = "Futura Std Book"
    if config.LOGO_PORTADA.exists():
        s.shapes.add_picture(str(config.LOGO_PORTADA), 4761876, 3683417, 2668248, 2668248)


def separador(deck: Deck, texto: str) -> None:
    s = deck.nueva(config.DISENO_SEPARADOR)
    _limpiar_marcadores(s, conservar_titulo=False)
    tam = 44 if len(texto) <= 40 else 36
    caja_texto(s, 442913, 1905507, 11306176, 3046988, texto, tam=tam,
               color=config.COLOR_MARCA_TEXTO, alinear=PP_ALIGN.CENTER,
               fuente="Futura Std Book", anclaje=MSO_ANCHOR.MIDDLE,
               nombre="Separador")


def subseparador(deck: Deck, texto: str) -> None:
    s = deck.nueva(config.DISENO_SUBSEPARADOR)
    _limpiar_marcadores(s, conservar_titulo=False)
    caja_texto(s, 968211, 543964, 10780877, 1077218, texto, tam=30, negrita=True,
               color=config.COLOR_MARCA_TEXTO, fuente="Futura Std Book",
               nombre="Subseparador")


# --------------------------------------------------------------------------- #
# I.a / II.a  Inversión (o TRP) y SOS (o SOV) por periodo
# --------------------------------------------------------------------------- #
def por_periodo(deck: Deck, ctx: Contexto, sec: Seccion, df: pd.DataFrame,
                metrica: str) -> None:
    es_trp = metrica == "trp"
    valor = sec.target if es_trp else "inv"
    base = df[df["es_tv"]] if es_trp else df
    tabla = metricas.por_periodo(base, ctx.per, sec.dim, valor)
    sel = metricas.seleccionar(tabla, ctx.per, ctx.op.top_grafico, sec.orden_fijo)
    sel = sel.loc[:, sel.fillna(0).sum() > 0]
    if sel.empty:
        ctx.avisos.append(f"{sec.nombre}: sin {'TRPs' if es_trp else 'inversión'}; "
                          "se omite la diapositiva.")
        return
    escala = 1 if es_trp else 1e6
    abs_ = sel / escala
    shares = metricas.participacion(sel)
    totales = abs_.sum(axis=1, min_count=1)
    entidades = list(sel.columns)
    colores = ctx.colores.para(entidades)
    cats = list(sel.index)

    s = deck.nueva(config.DISENO_CONTENIDO)
    _limpiar_marcadores(s, conservar_titulo=True)
    ent = sec.entidad
    titulo = f"TRPS Y SOV POR {ent}" if es_trp else f"INVERSIONES Y SOS POR {ent}"
    poner_titulo(s, titulo, _sub(sec, rango(ctx.per, ctx.per.anio_inicio)))
    kw = dict(umbral_etiqueta=0.05, tam=9, tam_etiquetas=9)
    columnas_apiladas(
        s, X0, Y0, ANCHO, ALTO, cats,
        [(e, abs_[e].tolist()) for e in entidades], colores,
        totales=[num(t) if pd.notna(t) else "" for t in totales],
        trazado=(0.07, 0.06, 0.41, 0.70),
        titulo_valor="TRPs" if es_trp else "CLP (millones de pesos)", **kw)
    columnas_apiladas(
        s, X0 + int(ANCHO * 0.52), Y0, int(ANCHO * 0.48), ALTO, cats,
        [(e, shares[e].tolist()) for e in entidades], colores, porcentaje=True,
        formato_etiqueta="0%", formato_valor="0%", leyenda=False,
        trazado=(0.07, 0.06, 0.90, 0.70), **kw)
    poner_fuente(s, _fuente_trp(sec) if es_trp else config.FUENTE_MILLONES)
    metrica_txt = "SOV" if es_trp else "SOS"
    _registrar(
        ctx, deck, sec, "TRP y SOV" if es_trp else "Inversión y SOS", titulo,
        _sub(sec, rango(ctx.per, ctx.per.anio_inicio)),
        metrica=(f"{sec.target} ({config.TARGETS_TRP.get(sec.target, '')})" if es_trp
                 else config.COL_VALOR),
        unidad_ppt=("TRPs (gráfico izquierdo) y SOV % (derecho)" if es_trp else
                    "Millones de pesos = valor / 1.000.000 (izquierdo) y SOS % (derecho)"),
        calculo=(f"Suma por periodo y serie. {metrica_txt} = valor de la serie / total del "
                 f"periodo. " + _regla_series(sec, ctx.op.top_grafico,
                                              "por SOS ponderado entre periodos")),
        reproducir=(f"Filtrar {sec.universo}{' y es_tv = VERDADERO' if es_trp else ''}. "
                    f"{_texto_periodos(ctx.per)} Sumar {valor if es_trp else config.COL_VALOR} "
                    f"por «{sec.dim}»."),
        bloques=[Bloque("Valores por periodo", tabla.T, entidades,
                        _asignacion(tabla.columns, entidades),
                        "TRP" if es_trp else "CLP", "columna")])


# --------------------------------------------------------------------------- #
# I.b / II.b  Tablas mensuales año anterior y actual
# --------------------------------------------------------------------------- #
ANCHOS_TABLA = [2150000] + [625000] * 12 + [950000, 700000]


def tabla_mensual(deck: Deck, ctx: Contexto, sec: Seccion, df: pd.DataFrame,
                  metrica: str) -> None:
    es_trp = metrica == "trp"
    valor = sec.target if es_trp else "inv"
    base = df[df["es_tv"]] if es_trp else df
    escala = 1 if es_trp else 1e3
    bloques = [metricas.mensual(base, a, sec.dim, valor, ctx.op.top_tabla, sec.orden_fijo)
               for a in (ctx.per.anio_anterior, ctx.per.anio_actual)]
    if all(b.total <= 0 for b in bloques):
        ctx.avisos.append(f"{sec.nombre}: sin {'TRPs' if es_trp else 'inversión'} mensual; "
                          "se omite la tabla.")
        return
    var, var_tot = metricas.variacion_mensual(base, ctx.per, valor)

    s = deck.nueva(config.DISENO_TABLA)
    _limpiar_marcadores(s, conservar_titulo=True)
    titulo = "TRPS MENSUALES" if es_trp else "INVERSIONES MENSUALES (000)"
    poner_titulo(s, titulo, _sub(sec, rango(ctx.per, ctx.per.anio_anterior)))

    etiqueta = "SOV" if es_trp else "SOS"
    dibujo = []
    for b in bloques:
        filas = []
        for ent, fila in b.valores.iterrows():
            tot = fila.sum(min_count=1)
            if pd.isna(tot) or tot == 0:
                continue
            filas.append((str(ent), [num(v / escala) if pd.notna(v) else "" for v in fila],
                          num(tot / escala), pct(tot / b.total if b.total else None)))
        total = ([num(v / escala) if pd.notna(v) else "" for v in b.total_mes],
                 num(b.total / escala), "100,0%" if b.total else "")
        dibujo.append(BloqueTabla(b.anio, filas, total))
    sep = 70000
    n_filas = sum(3 + len(b.filas) for b in dibujo) + 1
    alto = int(min(200000, (ALTO - len(dibujo) * sep) / n_filas))
    tam = 8 if alto >= 150000 else 7
    dibujar_tabla_mensual(s, X0, Y0, ANCHOS_TABLA, alto, bloques=dibujo,
                          etiqueta_share=etiqueta, variacion=var.tolist(),
                          variacion_total=var_tot, alto_separador=sep, tam=tam)
    poner_fuente(s, _fuente_trp(sec) if es_trp else config.FUENTE_MILES)
    if ctx.traza is not None:
        unidad = "TRP" if es_trp else "CLP"
        bloques_t = []
        for b in bloques:
            det = metricas.detalle_mensual(base, b.anio, sec.dim, valor).T
            bloques_t.append(Bloque(
                f"Año {b.anio}", det, list(b.valores.index),
                _asignacion(det.index, b.valores.index), unidad, "total",
                etiquetas_columnas=list(config.MESES_CORTOS)))
        comp, act_tot, base_tot = metricas.componentes_variacion(base, ctx.per, valor)
        _registrar(
            ctx, deck, sec, "TRP mensual" if es_trp else "Inversión mensual", titulo,
            _sub(sec, rango(ctx.per, ctx.per.anio_anterior)),
            metrica=(f"{sec.target} ({config.TARGETS_TRP.get(sec.target, '')})" if es_trp
                     else config.COL_VALOR),
            unidad_ppt="TRPs" if es_trp else "Miles de pesos = valor / 1.000",
            calculo=(f"Suma por mes y serie. {etiqueta} = total anual de la serie / total del "
                     f"año. " + _regla_series(sec, ctx.op.top_tabla,
                                              "por inversión del año (cada año por separado)")),
            reproducir=(f"Filtrar {sec.universo}{' y es_tv = VERDADERO' if es_trp else ''} y "
                        f"Año. Sumar {valor if es_trp else config.COL_VALOR} por «{sec.dim}» "
                        f"y Mes."),
            bloques=bloques_t,
            variacion=Variacion(comp, act_tot, base_tot, ctx.per.anio_actual,
                                ctx.per.anio_anterior, unidad))


# --------------------------------------------------------------------------- #
# I.c  SOS mensual por marca
# --------------------------------------------------------------------------- #
def sos_mensual(deck: Deck, ctx: Contexto, sec: Seccion, df: pd.DataFrame) -> None:
    t = metricas.share_mensual(df, ctx.per, sec.dim, "inv", ctx.op.top_grafico,
                               sec.orden_fijo)
    t = t.loc[:, t.fillna(0).sum() > 0]
    if t.empty:
        ctx.avisos.append(f"{sec.nombre}: sin inversión mensual; se omite SOS mensual.")
        return
    n_meses = len(t) - 2
    vacia = pd.DataFrame([[None] * len(t.columns)], columns=t.columns, index=[""])
    t = pd.concat([t.iloc[:n_meses], vacia, t.iloc[n_meses:]])
    shares = metricas.participacion(t)
    tot = t.sum(axis=1, min_count=1)
    entidades = list(t.columns)
    colores = ctx.colores.para(entidades)

    s = deck.nueva(config.DISENO_CONTENIDO)
    _limpiar_marcadores(s, conservar_titulo=True)
    poner_titulo(s, "SOS MENSUALES POR MARCA" if sec.entidad == "MARCA"
                 else f"SOS MENSUALES POR {sec.entidad}",
                 _sub(sec, rango(ctx.per, ctx.per.anio_anterior)))
    columnas_apiladas(
        s, X0, Y0, ANCHO, ALTO, list(t.index),
        [(e, shares[e].tolist()) for e in entidades], colores, porcentaje=True,
        formato_etiqueta="0%", formato_valor="0%",
        totales=[num(v / 1e3) if pd.notna(v) and v > 0 else "" for v in tot],
        umbral_etiqueta=0.12, tam=8, tam_etiquetas=7, rot_totales=-90,
        rot_categorias=-45, trazado=(0.05, 0.12, 0.93, 0.66),
        titulo_valor="SOS (miles de pesos sobre cada columna)")
    poner_fuente(s, config.FUENTE_MILES)
    if ctx.traza is not None:
        det = metricas.detalle_share_mensual(df, ctx.per, sec.dim, "inv").T
        _registrar(
            ctx, deck, sec, "SOS mensual",
            "SOS MENSUALES POR MARCA" if sec.entidad == "MARCA"
            else f"SOS MENSUALES POR {sec.entidad}",
            _sub(sec, rango(ctx.per, ctx.per.anio_anterior)),
            metrica=config.COL_VALOR,
            unidad_ppt="SOS % por columna; sobre cada columna, total en miles de pesos",
            calculo=("Suma por mes y serie; SOS = valor de la serie / total del mes. Las dos "
                     "últimas columnas son YTD. " + _regla_series(
                         sec, ctx.op.top_grafico, "(unión de los top de cada YTD)")),
            reproducir=(f"Filtrar {sec.universo}. Meses: Año y Mes. "
                        f"{_texto_periodos(ctx.per)} Sumar {config.COL_VALOR} por «{sec.dim}»."),
            bloques=[Bloque("Valores por mes y YTD", det, entidades,
                            _asignacion(det.index, entidades), "CLP", "columna")])


# --------------------------------------------------------------------------- #
# I.d  Mix de medios
# --------------------------------------------------------------------------- #
def mix_medios(deck: Deck, ctx: Contexto, sec: Seccion, df: pd.DataFrame) -> None:
    bloques = metricas.mix_medios(df, ctx.per, sec.dim, ctx.op.top_mix, sec.orden_fijo)
    if not bloques:
        ctx.avisos.append(f"{sec.nombre}: sin inversión; se omite mix de medios.")
        return
    filas, totales = [], []
    for k, b in enumerate(bloques):
        if k:
            filas.append(pd.Series({m: None for m in config.ORDEN_MEDIOS}, name=""))
            totales.append(None)
        for nombre, fila in b.shares.iterrows():
            filas.append(fila.rename(nombre))
            totales.append(b.totales[nombre])
    t = pd.DataFrame(filas)
    medios = [m for m in config.ORDEN_MEDIOS if t[m].fillna(0).sum() > 0]
    colores = {m: ctx.colores.fijos.get(m.upper(), "A5A5A5") for m in medios}

    s = deck.nueva(config.DISENO_CONTENIDO_LATERAL)
    _limpiar_marcadores(s, conservar_titulo=True)
    poner_titulo(s, "MIX DE MEDIOS", _sub(sec, f"TT {ctx.per.anio_anterior} Y "
                                          f"{rango(ctx.per)}"))
    columnas_apiladas(
        s, X0, Y0, ANCHO, ALTO, list(t.index),
        [(m, t[m].tolist()) for m in medios], colores, porcentaje=True,
        formato_etiqueta="0%", formato_valor="0%",
        totales=[num(v / 1e6) if v is not None and pd.notna(v) else "" for v in totales],
        umbral_etiqueta=0.06, tam=8, tam_etiquetas=8,
        rot_categorias=-45 if len(t) > 9 else 0, rot_totales=-90 if len(t) > 14 else 0,
        trazado=(0.05, 0.10, 0.93, 0.66))
    poner_fuente(s, config.FUENTE_MILLONES)
    if ctx.traza is not None:
        detalles = metricas.detalle_mix(df, ctx.per, sec.dim)
        bloques_t = []
        for b in bloques:
            visibles = [e for e in b.shares.index if e != b.etiqueta]
            det = detalles[b.etiqueta]
            bloques_t.append(Bloque(
                b.etiqueta, det, visibles, _asignacion(det.index, visibles), "CLP", "fila",
                etiqueta_total=b.etiqueta,
                nota=f"La fila «{b.etiqueta}» es la última barra del bloque en el gráfico."))
        _registrar(
            ctx, deck, sec, "Mix de medios", "MIX DE MEDIOS",
            _sub(sec, f"TT {ctx.per.anio_anterior} Y {rango(ctx.per)}"),
            metrica=f"{config.COL_VALOR} por medio_grp",
            unidad_ppt="% por medio en cada barra; sobre cada barra, total en millones de pesos",
            calculo=("Suma por serie y medio; % = valor del medio / total de la serie. "
                     + _regla_series(sec, ctx.op.top_mix, "por inversión del periodo")),
            reproducir=(f"Filtrar {sec.universo}. TT {ctx.per.anio_anterior}: Año = "
                        f"{ctx.per.anio_anterior}. YTD {ctx.per.anio_actual}: Año = "
                        f"{ctx.per.anio_actual}. Sumar {config.COL_VALOR} por «{sec.dim}» "
                        f"y «medio_grp»."),
            bloques=bloques_t)


# --------------------------------------------------------------------------- #
# III.  SOS (TV) vs SOV
# --------------------------------------------------------------------------- #
def sos_vs_sov(deck: Deck, ctx: Contexto, sec: Seccion, df: pd.DataFrame) -> None:
    r = metricas.sos_vs_sov(df, ctx.per, sec.dim, sec.target, ctx.op.top_sov,
                            sec.orden_fijo)
    if r is None:
        ctx.avisos.append(f"{sec.nombre}: sin TRPs en {ctx.per.anio_actual}; "
                          "se omite SOS vs SOV.")
        return
    colores = ctx.colores.para(r.entidades)
    s = deck.nueva(config.DISENO_CONTENIDO)
    _limpiar_marcadores(s, conservar_titulo=True)
    poner_titulo(s, "SOS VS SOV", _sub(sec, rango(ctx.per)))
    caja_texto(s, 560388, 1420000, 6000000, 330000,
               f"{num(r.trp_total)} TRPS / {num(r.inv_total)} CLP", tam=16,
               nombre="Totales TV")
    dona_doble(s, 2300000, 1750000, 7600000, 4650000, r.entidades,
               r.sos.tolist(), r.sov.tolist(), colores)
    caja_texto(s, 560388, 6150000, 6500000, 300000,
               "Anillo interior: SOS (inversión en TV)  ·  Anillo exterior: SOV (TRPs)",
               tam=11, color="595959", nombre="Nota anillos")
    poner_fuente(s, _fuente_trp(sec))
    if ctx.traza is not None:
        det = metricas.detalle_sos_sov(df, ctx.per, sec.dim, sec.target)
        asig = _asignacion(det.index, r.entidades)
        _registrar(
            ctx, deck, sec, "SOS vs SOV", "SOS VS SOV", _sub(sec, rango(ctx.per)),
            metrica=(f"{config.COL_VALOR} (anillo interior) y {sec.target} "
                     f"({config.TARGETS_TRP.get(sec.target, '')}, anillo exterior)"),
            unidad_ppt="% de cada anillo; arriba, TRPs y pesos totales",
            calculo=("Año actual, solo TV abierta y pagada. SOS = inversión en TV de la serie / "
                     "total; SOV = TRPs de la serie / total. " + (
                         "Series: todas, en orden fijo." if sec.orden_fijo is not None else
                         f"Series: top {ctx.op.top_sov} por TRPs (solo quienes tienen TRPs); "
                         "el resto de la inversión en TV va a OTRAS.")),
            reproducir=(f"Filtrar {sec.universo}, Año = {ctx.per.anio_actual} y es_tv = "
                        f"VERDADERO. Sumar {config.COL_VALOR} y {sec.target} por «{sec.dim}»."),
            notas=[f"Texto del PPT: {num(r.trp_total)} TRPS / {num(r.inv_total)} CLP"],
            bloques=[
                Bloque("Anillo interior: inversión en TV (SOS)", det[["inv"]], r.entidades,
                       asig, "CLP", "columna", etiquetas_columnas=["Inversión TV (CLP)"]),
                Bloque("Anillo exterior: TRPs (SOV)", det[["trp"]], r.entidades, asig, "TRP",
                       "columna", etiquetas_columnas=["TRPs"]),
            ])
