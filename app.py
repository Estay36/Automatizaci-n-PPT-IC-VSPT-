import streamlit as st
import pandas as pd
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches
from io import BytesIO
import matplotlib.ticker as ticker

# Configuración inicial de la página
st.set_page_config(page_title="Generador de Informe VSPT", layout="wide")
st.title("🍷 Generador Automático: Informe de Competencia VSPT")
st.markdown("Sube los archivos solicitados para procesar los datos y generar el PPTX automáticamente.")

# Subida de archivos
col1, col2, col3 = st.columns(3)
with col1:
    file_datos = st.file_uploader("1. Informe VSPT (Excel datos)", type=['xlsx', 'xls'])
with col2:
    file_cat = st.file_uploader("2. Categorías Vinos (Agrupaciones)", type=['xlsx', 'xls'])
with col3:
    file_ppt = st.file_uploader("3. Template Base (PPTX)", type=['pptx'])

def extraer_agrupaciones(archivo_categorias):
    """
    Lee la hoja 'Agrupaciones' identificando la fila 2 como el nombre de las compañías
    y mapeando las marcas que se encuentran hacia abajo en la misma columna.
    """
    df_agr = pd.read_excel(archivo_categorias, sheet_name='Agrupaciones', header=None)
    mapa_marcas = {}
    
    # Recorremos cada columna buscando nombres de compañías en la fila 2 (índice 2)
    for col in df_agr.columns:
        company = df_agr.iloc[2, col]
        if pd.notna(company):
            # Las marcas están desde la fila 3 hacia abajo
            marcas = df_agr.iloc[3:, col].dropna().tolist()
            for marca in marcas:
                marca_limpia = str(marca).strip().upper()
                mapa_marcas[marca_limpia] = str(company).strip().upper()
                
    return mapa_marcas

def add_slide_with_plot(prs, fig, slide_title):
    """
    Agrega una diapositiva al PPTX insertando un gráfico de Matplotlib generado.
    """
    # Intentar usar layout 'Title Only' (usualmente índice 5), si no, usar uno en blanco (6)
    try:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = slide_title
    except:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9), Inches(1))
        txBox.text_frame.text = slide_title
        
    # Guardar gráfico en memoria temporal (BytesIO)
    img_stream = BytesIO()
    fig.savefig(img_stream, format='png', bbox_inches='tight', dpi=150)
    img_stream.seek(0)
    
    # Insertar la imagen centrada
    slide.shapes.add_picture(img_stream, Inches(1), Inches(1.5), width=Inches(8))
    plt.close(fig)

if st.button("🚀 Procesar Datos y Generar Presentación", type="primary"):
    if file_datos and file_cat and file_ppt:
        with st.spinner("1. Extrayendo agrupaciones y limpiando datos..."):
            mapa_marcas = extraer_agrupaciones(file_cat)
            
            # Leer los datos principales
            df_datos = pd.read_excel(file_datos, sheet_name='datos')
            
            # Limpiar nombre de marcas y hacer cruce con mapa de agrupaciones
            df_datos['Marca_Clean'] = df_datos['Marcas reunidas'].astype(str).str.strip().str.upper()
            df_datos['Compañía_Agrupada'] = df_datos['Marca_Clean'].map(mapa_marcas).fillna('OTRAS VIÑAS')
            
            # Limpieza básica de Fechas / Valores
            df_datos['Año'] = pd.to_numeric(df_datos['Año'], errors='coerce')
            df_datos['ValorPublico'] = pd.to_numeric(df_datos['ValorPublico'], errors='coerce').fillna(0)

        with st.spinner("2. Generando gráficos e inyectándolos al Template PPTX..."):
            prs = Presentation(file_ppt)
            
            # ========================================================
            # DIAPOSITIVA 1: Inversión y SOS por Compañía (Total)
            # ========================================================
            inv_cia = df_datos.groupby('Compañía_Agrupada')['ValorPublico'].sum().sort_values(ascending=False)
            
            fig1, ax1 = plt.subplots(figsize=(10, 5))
            # Gráfico de Torta (SOS)
            wedges, texts, autotexts = ax1.pie(
                inv_cia, labels=inv_cia.index, autopct='%1.1f%%', 
                startangle=90, colors=plt.cm.tab20.colors
            )
            ax1.set_title("SOS Inversión por Compañía (Total Histórico)", fontsize=14, pad=20)
            add_slide_with_plot(prs, fig1, "Inversión y SOS por Compañía")

            # ========================================================
            # DIAPOSITIVA 2: Inversión Mensual 2025 y 2026 (Compañías)
            # ========================================================
            df_reciente = df_datos[df_datos['Año'].isin([2025, 2026])]
            if not df_reciente.empty:
                pivot_mensual = df_reciente.pivot_table(
                    index='Mes', columns=['Compañía_Agrupada', 'Año'], 
                    values='ValorPublico', aggfunc='sum'
                ).fillna(0)

                fig2, ax2 = plt.subplots(figsize=(12, 6))
                pivot_mensual.plot(kind='line', marker='o', ax=ax2)
                ax2.set_title("Inversión Mensual por Compañía (2025 vs 2026)", fontsize=14)
                ax2.set_ylabel("Valor Público (Inversión)")
                ax2.yaxis.set_major_formatter(ticker.FuncFormatter(lambda x, pos: f'${x:,.0f}'))
                ax2.grid(True, linestyle='--', alpha=0.6)
                plt.legend(bbox_to_anchor=(1.05, 1), loc='upper left')
                
                add_slide_with_plot(prs, fig2, "Inversión Mensual 2025 - 2026")

            # ========================================================
            # DIAPOSITIVA 3: Inversión por Categoría
            # ========================================================
            if 'Área' in df_datos.columns:
                inv_cat = df_datos.groupby('Área')['ValorPublico'].sum().sort_values(ascending=True)
                fig3, ax3 = plt.subplots(figsize=(10, 5))
                inv_cat.plot(kind='barh', ax=ax3, color='#8E244D') # Color corporativo VSPT
                ax3.set_title("Inversión Total por Categoría (Área)", fontsize=14)
                ax3.set_xlabel("Valor Público (CLP)")
                ax3.xaxis.set_major_formatter(ticker.FuncFormatter(lambda x, pos: f'${x:,.0f}'))
                add_slide_with_plot(prs, fig3, "Inversión por Categorías")

            # ========================================================
            # DIAPOSITIVA 4: Mix de Medios por Compañía
            # ========================================================
            if 'Medio' in df_datos.columns:
                mix_medios = df_datos.pivot_table(
                    index='Compañía_Agrupada', columns='Medio', values='ValorPublico', aggfunc='sum'
                ).fillna(0)
                
                # Convertir a porcentajes (100% stacked bar)
                mix_medios_pct = mix_medios.div(mix_medios.sum(axis=1), axis=0) * 100
                
                fig4, ax4 = plt.subplots(figsize=(12, 6))
                mix_medios_pct.plot(kind='bar', stacked=True, ax=ax4, colormap='Spectral')
                ax4.set_title("Mix de Medios por Compañía (%)", fontsize=14)
                ax4.set_ylabel("Porcentaje de Inversión (%)")
                ax4.legend(bbox_to_anchor=(1.05, 1), loc='upper left')
                plt.xticks(rotation=45)
                add_slide_with_plot(prs, fig4, "Mix de Medios por Compañía")

            # Guardar el documento PPTX final en memoria
            output_pptx = BytesIO()
            prs.save(output_pptx)
            output_pptx.seek(0)
            
        st.success("¡Informe PPTX generado con éxito! Las métricas clave se inyectaron basadas en tus agrupaciones de Excel.")
        
        st.download_button(
            label="📥 Descargar Informe Competencia VSPT (PPTX)",
            data=output_pptx,
            file_name="Informe_Competencia_Automatizado.pptx",
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )
    else:
        st.warning("⚠️ Por favor, sube los 3 archivos solicitados antes de generar el reporte.")
