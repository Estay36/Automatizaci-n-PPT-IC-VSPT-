"""Pruebas con datos sintéticos (la base real es confidencial y no se versiona).

Ejecutar:  python -m pytest -q
"""
from __future__ import annotations

import io
import itertools
from datetime import date

import openpyxl
import pandas as pd
import pytest
from pptx import Presentation

from informe import config
from informe.agrupaciones import clave, leer_agrupaciones
from informe.datos import fecha_corte_desde_nombre, preparar
from informe.diapositivas import Opciones
from informe.reporte import generar

# (marca, área, empresa)
MARCAS = [
    ("GATO", "VINO MASIVO", "VIÑA SAN PEDRO - TARAPACÁ"),
    ("GATO", "VINO-COOLER-SANGRÍA", "VIÑA SAN PEDRO - TARAPACÁ"),
    ("DONNALUNA", "VINO-COOLER-SANGRÍA", "VIÑA SAN PEDRO - TARAPACÁ"),
    ("MISIONES DE RENGO", "VINO PREMIUM", "VIÑA SAN PEDRO - TARAPACÁ"),
    ("CASILLERO DEL DIABLO", "VINO PREMIUM", "VIÑA CONCHA Y TORO"),
    ("CASILLERO DEL DIABLO", "ESPUMANTES", "VIÑA CONCHA Y TORO"),
    ("120", "VINO MASIVO", "VIÑA SANTA RITA"),
    ("120 SPRITZ", "VINO-COOLER-SANGRÍA", "VIÑA SANTA RITA"),
    ("UNDURRAGA", "VINO MASIVO", "VIÑA UNDURRAGA"),
    ("UNDURRAGA", "ESPUMANTES", "VIÑA UNDURRAGA"),
    ("VALDIVIESO", "ESPUMANTES", "LICORES MITJANS"),
    ("MANQUEHUITO", "VINO-COOLER-SANGRÍA", "VIÑA SAN PEDRO - TARAPACÁ"),
    ("STONES", "CERVEZAS", "C.C.U. S.A."),
    ("MISTRAL ICE", "ICE", "COMPAÑÍA PISQUERA DE CHILE S.A."),
    ("MARCA NUEVA", "VINO PREMIUM", "VIÑA DESCONOCIDA"),
]
MEDIOS = ["TV Abierta", "Radio", "Vía Pública", "Prensa"]


@pytest.fixture(scope="module")
def agrupaciones():
    return leer_agrupaciones(config.AGRUPACIONES_DEFECTO)


@pytest.fixture(scope="module")
def base() -> pd.DataFrame:
    filas = []
    for (marca, area, emp), anio, mes, medio in itertools.product(
            MARCAS, (2024, 2025, 2026), range(1, 13), MEDIOS):
        if anio == 2026 and mes > 10:
            continue
        semana = min(52, (mes - 1) * 4 + 2)
        es_tv = medio == "TV Abierta"
        filas.append({
            "Área": area, "Empresa_Clasif": emp, "Marcas reunidas": marca,
            "Año": anio, "Mes": mes, "Semana": semana, "Medio": medio,
            "ValorPublico": 1_000_000 * (1 + len(marca) % 5) * (2 if es_tv else 1),
            "HM25_49MB_Pond": 10.0 if es_tv else 0.0,
            "HM35_49A_Pond": 8.0 if es_tv else 0.0,
            "HM18_49AMB_Pond": 12.0 if es_tv else 0.0,
            "Periodo": "COMPARABLE" if (anio == 2026 or mes <= 9) else "NO COMPARABLE",
        })
    return pd.DataFrame(filas)


def test_clave_normaliza():
    assert clave("Royal Guard (Sabores)*") == "ROYALGUARD"
    assert clave("B-MIX") == clave("BMIX")
    assert clave("ÉPICA") == clave("EPICA")


def test_fecha_desde_nombre():
    assert fecha_corte_desde_nombre("Informe_VSPT_2022-02oct2026.xlsx") == date(2026, 10, 2)
    assert fecha_corte_desde_nombre("sin_fecha.xlsx") is None


def test_clasificacion(base, agrupaciones):
    df, per, diag = preparar(base, agrupaciones, date(2026, 10, 2))
    cat = (df.drop_duplicates(["Marcas reunidas", "Área"])
             .set_index(["Marcas reunidas", "Área"])["categoria"])
    assert cat[("DONNALUNA", "VINO-COOLER-SANGRÍA")] == "BOTELLAS FINAS"
    assert cat[("UNDURRAGA", "VINO MASIVO")] == "MASIVOS"
    assert cat[("UNDURRAGA", "ESPUMANTES")] == "ESPUMANTES"
    assert cat[("CASILLERO DEL DIABLO", "ESPUMANTES")] == "ESPUMANTES"
    assert cat[("120 SPRITZ", "VINO-COOLER-SANGRÍA")] == "MASIVOS"
    # Fusión 120 SPRITZ -> 120
    assert set(df.loc[df["Marcas reunidas"] == "120 SPRITZ", "marca"]) == {"120"}
    # GATO LATA es BLYS vinos; GATO botella no.
    blys = df[df["blys_vinos"]]["marca_blys"].unique()
    assert "GATO LATA" in blys and "GATO" not in blys
    # 120 SPRITZ está en "tentativo a futuro": fuera de BLYS.
    assert not df.loc[df["Marcas reunidas"] == "120 SPRITZ", "blys_total"].any()
    assert df.loc[df["Marcas reunidas"] == "STONES", "blys_total"].all()
    # Compañía
    comp = df.drop_duplicates("Marcas reunidas").set_index("Marcas reunidas")["compania"]
    assert comp["GATO"] == "VSPT" and comp["120"] == "SANTA RITA"
    assert comp["VALDIVIESO"] == config.COMPANIA_OTRAS
    assert "MARCA NUEVA" in set(diag.sin_categoria["Marcas reunidas"])
    assert per.etiquetas == ["TT 2024", "TT 2025", "YTD 2025", "YTD 2026"]


def test_genera_pptx(base, agrupaciones):
    res = generar(base, agrupaciones, Opciones(fecha_corte=date(2026, 10, 2)),
                  trazabilidad=False)
    assert res.xlsx is None
    prs = Presentation(io.BytesIO(res.pptx))
    assert len(prs.slides) > 40
    titulos = [s.shapes.title.text_frame.text for s in prs.slides
               if s.shapes.title is not None]
    assert any(t.startswith("INVERSIONES Y SOS POR COMPAÑÍA") for t in titulos)
    assert any(t.startswith("SOS VS SOV") for t in titulos)


def test_trazabilidad_cuadra_con_pptx(base, agrupaciones):
    """Cada diapositiva con datos tiene su hoja y la numeración calza con el PPT."""
    res = generar(base, agrupaciones, Opciones(fecha_corte=date(2026, 10, 2)),
                  incluir_base=True)
    prs = Presentation(io.BytesIO(res.pptx))
    wb = openpyxl.load_workbook(io.BytesIO(res.xlsx))
    assert {"LEEME", "Índice", "Clasificación", "Diagnóstico", "Base"} <= set(wb.sheetnames)
    indice = wb["Índice"]
    filas = [r for r in indice.iter_rows(min_row=4, values_only=True) if isinstance(r[0], int)]
    assert len(filas) == len(prs.slides)  # una fila por diapositiva
    for n, _sec, tipo, titulo, *_resto, hoja in filas:
        slide = prs.slides[n - 1]
        if hoja not in (None, "—"):
            assert hoja in wb.sheetnames
            assert slide.shapes.title.text_frame.text.startswith(titulo)
    # La tabla A de "Inversión por compañía" suma lo mismo que la base para YTD actual.
    hoja = wb[next(s for s in wb.sheetnames if s.endswith("Compañía Inv"))]
    filas_h = list(hoja.iter_rows(values_only=True))
    i_enc = next(i for i, r in enumerate(filas_h) if r[0] == "Entidad")
    col = filas_h[i_enc].index("YTD 2026")
    suma = 0.0
    for r in filas_h[i_enc + 1:]:
        if r[0] is None:
            break
        suma += r[col] or 0
    esperado = base[base["Área"].isin(config.AREAS_VINO) & (base["Año"] == 2026)][
        "ValorPublico"].sum()
    assert suma == pytest.approx(esperado)
